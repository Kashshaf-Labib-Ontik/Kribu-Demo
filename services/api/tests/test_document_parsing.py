from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from app.documents.parsing import ParseQuarantinedError, parse_document


def _docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_heading("Milestone review", level=1)
    document.add_paragraph("The permit is still pending.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Programs"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Milestones"
    sheet["A1"] = "Milestone"
    sheet["B2"] = "Delayed"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    stream = b"BT /F1 12 Tf 72 720 Td (Milestone M-17 is delayed.) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, start=1):
        offsets.append(len(result))
        result.extend(f"{index} 0 obj\n".encode())
        result.extend(value)
        result.extend(b"\nendobj\n")
    startxref = len(result)
    result.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    result.extend(b"".join(f"{offset:010} 00000 n \n".encode() for offset in offsets[1:]))
    trailer = (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{startxref}\n%%EOF\n"
    )
    result.extend(trailer.encode())
    return bytes(result)


@pytest.mark.parametrize(
    ("filename", "data", "expected_locator"),
    [
        pytest.param(
            "report.csv",
            b"period,status\nMarch,Delayed\n",
            {"kind": "csv_cell", "row": 2},
            id="csv",
        ),
        pytest.param(
            "notes.md",
            b"# Risks\nPermit delay remains unresolved.\n",
            {"kind": "markdown_line", "line": 2, "heading_path": ["Risks"]},
            id="markdown",
        ),
        pytest.param(
            "history.docx", _docx_bytes(), {"kind": "docx_paragraph", "paragraph": 1}, id="docx"
        ),
        pytest.param(
            "status.xlsx",
            _xlsx_bytes(),
            {"kind": "xlsx_cell", "sheet": "Milestones", "cell": "A1"},
            id="xlsx",
        ),
        pytest.param("report.pdf", _pdf_bytes(), {"kind": "pdf_page", "page": 1}, id="pdf"),
    ],
)
def test_parsers_preserve_canonical_source_locators(
    filename: str, data: bytes, expected_locator: dict[str, object]
) -> None:
    result = parse_document(data=data, filename=filename, max_pdf_pages=250, max_content_units=100)

    assert result.blocks
    assert any(
        all(block.locator.get(key) == value for key, value in expected_locator.items())
        for block in result.blocks
    )
    assert [block.ordinal for block in result.blocks] == list(range(1, len(result.blocks) + 1))


def test_content_unit_limit_quarantines_without_a_partial_result() -> None:
    with pytest.raises(ParseQuarantinedError, match="content_unit_limit_exceeded"):
        parse_document(
            data=b"first\nsecond\n",
            filename="report.txt",
            max_pdf_pages=250,
            max_content_units=1,
        )


def test_invalid_ooxml_is_quarantined() -> None:
    with pytest.raises(ParseQuarantinedError, match="invalid_ooxml_archive"):
        parse_document(
            data=b"not a zip",
            filename="bad.docx",
            max_pdf_pages=250,
            max_content_units=100,
        )
