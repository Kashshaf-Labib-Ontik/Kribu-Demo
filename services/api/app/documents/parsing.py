"""Deterministic, CPU-only parsers that preserve source-specific citation locators."""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from importlib.metadata import version
from pathlib import PurePosixPath

from docling_parse.pdf_parser import DoclingPdfParser
from docx import Document as DocxDocument
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pypdf import PdfReader

from app.documents.validation import extension_of

NORMALIZATION_VERSION = "v1"
MAX_OOXML_UNCOMPRESSED_BYTES = 100 * 1024 * 1024


class ParseQuarantinedError(Exception):
    """A safe, stable parsing failure that must not publish partial content."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class ParsedBlock:
    ordinal: int
    block_type: str
    text: str
    locator: dict[str, object]


@dataclass(frozen=True)
class ParseResult:
    parser_name: str
    parser_version: str
    blocks: list[ParsedBlock]


def parse_document(
    *,
    data: bytes,
    filename: str,
    max_pdf_pages: int,
    max_content_units: int,
) -> ParseResult:
    """Parse one already-validated immutable original or raise a safe quarantine error."""
    parser: dict[str, Callable[[], ParseResult]] = {
        ".pdf": lambda: _parse_pdf(data, max_pdf_pages=max_pdf_pages),
        ".docx": lambda: _parse_docx(data, max_content_units=max_content_units),
        ".xlsx": lambda: _parse_xlsx(data, max_content_units=max_content_units),
        ".csv": lambda: _parse_csv(data, max_content_units=max_content_units),
        ".txt": lambda: _parse_text(data, markdown=False, max_content_units=max_content_units),
        ".md": lambda: _parse_text(data, markdown=True, max_content_units=max_content_units),
    }
    try:
        result = parser[extension_of(filename)]()
    except KeyError as exc:
        raise ParseQuarantinedError("unsupported_parser") from exc
    except ParseQuarantinedError:
        raise
    except Exception as exc:  # parser libraries intentionally never leak their errors to callers
        raise ParseQuarantinedError("parse_failed") from exc
    if not result.blocks:
        raise ParseQuarantinedError("no_extractable_content")
    if len(result.blocks) > max_content_units:
        raise ParseQuarantinedError("content_unit_limit_exceeded")
    return result


def _parse_pdf(data: bytes, *, max_pdf_pages: int) -> ParseResult:
    """Use Docling Parse plus pypdf page locators for exact citations.

    Docling Parse is deliberately kept CPU-only and OCR-free for the demo.  ``pypdf`` supplies
    the deterministic page boundary that the preview and future citation contract requires.
    """
    try:
        # This native Docling component validates and decodes programmatic PDF structure without
        # pulling Docling's layout/OCR/VLM models.  Text is read page-by-page below so each block
        # has an exact, stable page locator.
        docling_document = DoclingPdfParser().load(path_or_stream=io.BytesIO(data))
        list(docling_document.iterate_pages())
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise ParseQuarantinedError("encrypted_pdf")
        if len(reader.pages) > max_pdf_pages:
            raise ParseQuarantinedError("pdf_page_limit_exceeded")
        blocks = [
            ParsedBlock(
                ordinal=index,
                block_type="page_text",
                text=text.strip(),
                locator={"kind": "pdf_page", "page": index},
            )
            for index, page in enumerate(reader.pages, start=1)
            if (text := (page.extract_text() or "")).strip()
        ]
    except ParseQuarantinedError:
        raise
    except Exception as exc:
        raise ParseQuarantinedError("invalid_pdf") from exc
    return ParseResult("docling_parse_cpu", _package_version("docling-parse"), blocks)


def _parse_docx(data: bytes, *, max_content_units: int) -> ParseResult:
    _check_ooxml_expansion(data)
    document = DocxDocument(io.BytesIO(data))
    blocks: list[ParsedBlock] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            style_name = paragraph.style.name if paragraph.style is not None else ""
            _append(
                blocks,
                "heading" if style_name.lower().startswith("heading") else "paragraph",
                text,
                {"kind": "docx_paragraph", "paragraph": index, "style": style_name},
                max_content_units,
            )
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            for column_index, cell in enumerate(row.cells, start=1):
                text = cell.text.strip()
                if text:
                    _append(
                        blocks,
                        "table_cell",
                        text,
                        {
                            "kind": "docx_table_cell",
                            "table": table_index,
                            "row": row_index,
                            "column": column_index,
                        },
                        max_content_units,
                    )
    return ParseResult("python_docx", _package_version("python-docx"), blocks)


def _parse_xlsx(data: bytes, *, max_content_units: int) -> ParseResult:
    _check_ooxml_expansion(data)
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    blocks: list[ParsedBlock] = []
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is None or str(cell.value).strip() == "":
                        continue
                    _append(
                        blocks,
                        "spreadsheet_cell",
                        str(cell.value),
                        {"kind": "xlsx_cell", "sheet": worksheet.title, "cell": cell.coordinate},
                        max_content_units,
                    )
    finally:
        workbook.close()
    return ParseResult("openpyxl", _package_version("openpyxl"), blocks)


def _parse_csv(data: bytes, *, max_content_units: int) -> ParseResult:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ParseQuarantinedError("invalid_utf8") from exc
    blocks: list[ParsedBlock] = []
    for row_index, row in enumerate(csv.reader(io.StringIO(text)), start=1):
        for column_index, value in enumerate(row, start=1):
            if value.strip():
                _append(
                    blocks,
                    "csv_cell",
                    value.strip(),
                    {"kind": "csv_cell", "row": row_index, "column": column_index},
                    max_content_units,
                )
    return ParseResult("python_csv", "stdlib", blocks)


def _parse_text(data: bytes, *, markdown: bool, max_content_units: int) -> ParseResult:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ParseQuarantinedError("invalid_utf8") from exc
    blocks: list[ParsedBlock] = []
    headings: list[str] = []
    for line_number, line in enumerate(text.splitlines(), start=1):
        value = line.strip()
        if not value:
            continue
        block_type = "line"
        if markdown and value.startswith("#"):
            level = len(value) - len(value.lstrip("#"))
            heading = value[level:].strip()
            if heading:
                headings = headings[: level - 1] + [heading]
                block_type = "heading"
        _append(
            blocks,
            block_type,
            value,
            {
                "kind": "markdown_line" if markdown else "text_line",
                "line": line_number,
                "heading_path": headings.copy(),
            },
            max_content_units,
        )
    return ParseResult("markdown" if markdown else "text", "stdlib", blocks)


def _append(
    blocks: list[ParsedBlock],
    block_type: str,
    text: str,
    locator: dict[str, object],
    limit: int,
) -> None:
    if len(blocks) >= limit:
        raise ParseQuarantinedError("content_unit_limit_exceeded")
    blocks.append(ParsedBlock(len(blocks) + 1, block_type, text, locator))


def _check_ooxml_expansion(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(item.file_size for item in archive.infolist())
            if total > MAX_OOXML_UNCOMPRESSED_BYTES:
                raise ParseQuarantinedError("archive_expansion_limit_exceeded")
            if any(
                PurePosixPath(item.filename).is_absolute()
                or ".." in PurePosixPath(item.filename).parts
                for item in archive.infolist()
            ):
                raise ParseQuarantinedError("invalid_ooxml_archive")
    except ParseQuarantinedError:
        raise
    except (OSError, RuntimeError, ValueError, zipfile.BadZipFile) as exc:
        raise ParseQuarantinedError("invalid_ooxml_archive") from exc


def _package_version(package: str) -> str:
    try:
        return version(package)
    except Exception:
        return "unknown"
